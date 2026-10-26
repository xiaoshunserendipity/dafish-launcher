"""把 docs/用户指南.md 转成 docx（版式对齐 操作手册.docx）。

用法：
    python scripts/make_guide.py                     # 覆盖 docs/用户指南.docx
    python scripts/make_guide.py 输出路径.docx        # 指定输出（例如放进发布包）

依赖 python-docx：pip install python-docx
"""

from __future__ import annotations

import os
import re
import sys
import time
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "用户指南.md"
OUT = (
    Path(sys.argv[1]).resolve()
    if len(sys.argv) > 1
    else ROOT / "docs" / "用户指南.docx"
)

CN_FONT = "微软雅黑"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x59, 0x59, 0x59)


def app_version() -> str:
    """从 app/__init__.py 读版本号，避免封面写死后和代码不一致。"""
    try:
        text = (ROOT / "app" / "__init__.py").read_text(encoding="utf-8-sig")
        match = re.search(r'__version__\s*=\s*["\']([^"\']+)', text)
        return match.group(1) if match else "?"
    except Exception:
        return "?"


def style_run(run, size=10.5, bold=False, color=None, mono=False):
    font_name = MONO_FONT if mono else CN_FONT
    run.font.name = font_name
    run.font.size = Pt(size - 0.5 if mono else size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text, size=10.5, color=None, bold=False):
    """处理 **加粗** 和 `代码`（支持嵌套，如 **`.lnk` 说明**）。"""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            add_inline(paragraph, piece[2:-2], size=size, color=color, bold=True)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            style_run(
                paragraph.add_run(piece[1:-1]),
                size=size,
                bold=bold,
                color=color,
                mono=True,
            )
        else:
            style_run(paragraph.add_run(piece), size=size, bold=bold, color=color)


def para(doc, text="", size=10.5, bold=False, color=None, align=None, space_after=5, mono=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.4
    if align is not None:
        p.alignment = align
    if text:
        if mono:
            style_run(p.add_run(text), size=size, bold=bold, color=color, mono=True)
        else:
            add_inline(p, text, size=size, color=color)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7 + 0.6 * level)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.35
    style_run(p.add_run(("• " if level == 0 else "– ")), size=10.5)
    add_inline(p, text, size=10.5)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    size = 16 if level == 1 else (13 if level == 2 else 11.5)
    style_run(h.add_run(text.replace("`", "")), size=size, bold=True, color=ACCENT)
    return h


def _shade(color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    return shd


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = t.rows[0].cells[index]
        cell.text = ""
        style_run(
            cell.paragraphs[0].add_run(header),
            size=9.5,
            bold=True,
            color=RGBColor(0xFF, 0xFF, 0xFF),
        )
        cell._tc.get_or_add_tcPr().append(_shade("1F4E79"))
    for row in rows:
        cells = t.add_row().cells
        for index, value in enumerate(row):
            if index >= len(cells):
                break
            cells[index].text = ""
            add_inline(cells[index].paragraphs[0], value, size=9.5)
    if widths:
        for row in t.rows:
            for index, width in enumerate(widths):
                if index < len(row.cells):
                    row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def split_row(line: str) -> list[str]:
    body = line.strip().strip("|")
    return [cell.strip() for cell in body.split("|")]


def is_separator(line: str) -> bool:
    cells = split_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c)


def build() -> None:
    lines = SRC.read_text(encoding="utf-8").splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    # ---------- 封面 ----------
    for _ in range(7):
        doc.add_paragraph()
    para(doc, "大肥鱼启动器", size=30, bold=True, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "用 户 指 南", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=26)
    para(doc, f"适用版本：v{app_version()}", size=12, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "更新日期：2026-10-26", size=12, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=54)
    para(doc, "适用系统：Windows 10 / 11 64 位", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "无需安装 Python 或任何依赖，双击 exe 即可运行", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # ---------- 正文 ----------
    index = 0
    total = len(lines)
    first_h1_skipped = False

    while index < total:
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            index += 1
            block: list[str] = []
            while index < total and not lines[index].strip().startswith("```"):
                block.append(lines[index].rstrip())
                index += 1
            index += 1
            for code_line in block:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                style_run(p.add_run(code_line or " "), size=9, mono=True, color=RGBColor(0x33, 0x33, 0x33))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 表格
        if stripped.startswith("|") and index + 1 < total and is_separator(lines[index + 1]):
            headers = split_row(stripped)
            index += 2
            rows: list[list[str]] = []
            while index < total and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index].strip()))
                index += 1
            cols = len(headers)
            widths = [16.6 / cols] * cols
            table(doc, headers, rows, widths)
            continue

        if not stripped:
            index += 1
            continue

        if stripped == "---":
            index += 1
            continue

        if stripped.startswith("### "):
            heading(doc, stripped[4:], 3)
        elif stripped.startswith("## "):
            heading(doc, stripped[3:], 2)
        elif stripped.startswith("# "):
            title = stripped[2:]
            # 文档标题已在封面出现，正文里跳过第一个 H1
            if not first_h1_skipped:
                first_h1_skipped = True
            else:
                heading(doc, title, 1)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.35
            add_inline(p, stripped[2:], size=10, color=GRAY)
        elif stripped.startswith("- "):
            bullet(doc, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            bullet(doc, stripped)
        else:
            para(doc, stripped)

        index += 1

    para(doc, "—— 指南结束 ——", color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

    # 先写临时文件再替换：docx 可能正被预览面板占用
    tmp_out = OUT.with_name(OUT.stem + ".new" + OUT.suffix)
    doc.save(str(tmp_out))
    last_error = None
    for _ in range(20):
        try:
            os.replace(tmp_out, OUT)
            print("saved:", OUT)
            return
        except PermissionError as exc:
            last_error = exc
            time.sleep(0.5)
    raise SystemExit(f"无法覆盖 {OUT}（可能被其他程序占用）：{last_error}\n新文件在：{tmp_out}")


if __name__ == "__main__":
    build()
